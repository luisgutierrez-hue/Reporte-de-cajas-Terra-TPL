import streamlit as st
import pandas as pd
import io
import re
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# --- CONFIGURACIÓN DE LA PÁGINA ---
st.set_page_config(page_title="Generador de Reportes de Caja", layout="centered")

st.title("📊 Generador de Reportes: Cajas Terra")
st.write("Sube el reporte original de Excel para generar automáticamente el reporte filtrado y formateado.")

archivo_subido = st.file_uploader("Sube el archivo Excel original (ej. tpl_Id_de_caja_...)", type=["xlsx"])

if archivo_subido is not None:
    try:
        # 1. EXTRAER LA FECHA DEL NOMBRE DEL ARCHIVO ORIGINAL
        match_fecha = re.search(r"(\d{2}_\d{2}_\d{4})", archivo_subido.name)
        if match_fecha:
            fecha_reporte = match_fecha.group(1)
        else:
            fecha_reporte = "Fecha_Desconocida"
            
        nombre_archivo_salida = f"Reportes de cajas {fecha_reporte} Terra.xlsx"

        st.info("Procesando los datos, por favor espera...")
        
        # 2. ENCONTRAR AUTOMÁTICAMENTE LA FILA DE LOS ENCABEZADOS
        df_raw = pd.read_excel(archivo_subido, sheet_name="Reporte de Id de caja", header=None)
        
        try:
            fila_encabezado = df_raw[df_raw.eq('Nombre de Usuario').any(axis=1)].index[0]
        except IndexError:
            st.error("❌ No se encontró la columna 'Nombre de Usuario' en el archivo. Verifica que sea el formato correcto.")
            st.stop()
            
        df = pd.read_excel(archivo_subido, sheet_name="Reporte de Id de caja", header=fila_encabezado)

        # 3. FILTRAR USUARIOS EXCLUIDOS
        usuarios_excluidos = ['nherrer', 'rafa', 'sebag', 'rodrigo', 'admin']
        df = df[~df['Nombre de Usuario'].str.lower().isin([u.lower() for u in usuarios_excluidos])].copy()

        # 4. LIMPIEZA Y TRANSFORMACIÓN DE DATOS
        df['Fecha y hora apertura caja'] = pd.to_datetime(df['Fecha y hora apertura caja'], errors='coerce').dt.strftime('%d/%m/%Y')

        def limpiar_moneda(valor):
            if pd.isna(valor):
                return 0.0
            if isinstance(valor, str):
                valor = valor.replace('$', '').replace('.', '').replace(',', '.').strip()
            try:
                return float(valor)
            except ValueError:
                return 0.0

        # Función auxiliar para encontrar columnas por su Letra de Excel (N, AI, AB, etc.)
        def obtener_columna_por_letra(dataframe, letra):
            num = 0
            for c in letra.upper():
                num = num * 26 + (ord(c) - ord('A')) + 1
            idx = num - 1 # Convertir a índice 0-based para Python
            
            # Si el archivo tiene suficientes columnas, extrae la que pedimos
            if idx < len(dataframe.columns):
                return dataframe.iloc[:, idx]
            else:
                # Si la columna no existe en el Excel, devuelve ceros para no romper la suma
                return pd.Series([0.0] * len(dataframe), index=dataframe.index)

        # --- NUEVAS REGLAS DE NEGOCIO ---
        # "Total efectivo depósito" = Columna N + Columna AI
        col_N = obtener_columna_por_letra(df, 'N').apply(limpiar_moneda)
        col_AI = obtener_columna_por_letra(df, 'AI').apply(limpiar_moneda)
        df['Total efectivo depósito_calc'] = col_N + col_AI

        # "Tarjetas" = Columna AB + Columna AW
        col_AB = obtener_columna_por_letra(df, 'AB').apply(limpiar_moneda)
        col_AW = obtener_columna_por_letra(df, 'AW').apply(limpiar_moneda)
        df['Tarjetas_calc'] = col_AB + col_AW

        # 5. CÁLCULO DEL TOTAL FINAL
        df['Total_calc'] = df['Total efectivo depósito_calc'] + df['Tarjetas_calc']

        # 6. SELECCIONAR Y RENOMBRAR LAS COLUMNAS PARA EL REPORTE FINAL
        df_final = pd.DataFrame({
            'Fecha': df['Fecha y hora apertura caja'],
            'Código': df['Nombre de Usuario'],
            'Nombre y Documento': df['Nombre y Documento'],
            'Agencia': df['Agencia'],
            'Total efectivo depósito': df['Total efectivo depósito_calc'], # Columna E
            'Tarjetas': df['Tarjetas_calc'],                               # Columna F
            'Total': df['Total_calc']                                      # Columna G
        })

        # 7. GENERAR EL EXCEL CON DISEÑO PROFESIONAL
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            df_final.to_excel(writer, index=False, sheet_name='Reporte Terra')
            
            workbook = writer.book
            worksheet = workbook['Reporte Terra']
            
            color_encabezado = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            fuente_encabezado = Font(color="FFFFFF", bold=True)
            color_cebra = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            formato_moneda = '#,##0' 
            
            for cell in worksheet[1]:
                cell.fill = color_encabezado
                cell.font = fuente_encabezado
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row), start=2):
                if row_idx % 2 == 0:
                    for cell in row:
                        cell.fill = color_cebra
                
                # Aplicar formato de moneda
                row[4].number_format = formato_moneda # E: Total efectivo
                row[5].number_format = formato_moneda # F: Tarjetas
                row[6].number_format = formato_moneda # G: Total
            
            for col in worksheet.columns:
                max_length = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                worksheet.column_dimensions[col_letter].width = max_length + 3

        # 8. BOTÓN DE DESCARGA
        st.success("✅ ¡El reporte se ha generado correctamente con las nuevas sumas!")
        st.download_button(
            label="📥 Descargar Reporte Generado",
            data=buffer.getvalue(),
            file_name=nombre_archivo_salida,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        st.error(f"❌ Ocurrió un error inesperado. Detalle técnico: {e}")
